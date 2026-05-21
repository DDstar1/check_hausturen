"""Step 1 — WooCommerce credentials + Excel file + table selection."""

import os
import openpyxl

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QFileDialog, QComboBox, QGroupBox, QMessageBox, QSizePolicy, QLineEdit,
)


class Step1(QWidget):
    def __init__(self):
        super().__init__()
        self.wb = None
        self.file_path = ""
        self._table_map = {}
        self._build()

    def _build(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(16)
        layout.setContentsMargins(40, 28, 40, 28)

        title = QLabel("Step 1 — Connect & Select Excel File")
        title.setObjectName("step_title")
        layout.addWidget(title)

        sub = QLabel(
            "Enter your WooCommerce store credentials, "
            "then choose your masterlist Excel file and table."
        )
        sub.setObjectName("step_sub")
        sub.setWordWrap(True)
        layout.addWidget(sub)
        layout.addSpacing(8)

        # ── Credentials ──────────────────────────────────────────────────────
        sg = QGroupBox("WooCommerce Store")
        sgl = QVBoxLayout(sg)

        row1 = QHBoxLayout()
        row1.addWidget(QLabel("URL:"))
        self.url_edit = QLineEdit()
        self.url_edit.setPlaceholderText("https://…")
        row1.addWidget(self.url_edit, stretch=1)
        sgl.addLayout(row1)

        row2 = QHBoxLayout()
        row2.addWidget(QLabel("Consumer Key:"))
        self.key_edit = QLineEdit()
        self.key_edit.setPlaceholderText("ck_…")
        row2.addWidget(self.key_edit, stretch=1)
        row2.addSpacing(16)
        row2.addWidget(QLabel("Consumer Secret:"))
        self.secret_edit = QLineEdit()
        self.secret_edit.setPlaceholderText("cs_…")
        self.secret_edit.setEchoMode(QLineEdit.EchoMode.Password)
        row2.addWidget(self.secret_edit, stretch=1)
        sgl.addLayout(row2)

        layout.addWidget(sg)

        # ── Excel file ───────────────────────────────────────────────────────
        fg = QGroupBox("Excel File")
        fgl = QHBoxLayout(fg)
        self.file_label = QLabel("No file selected")
        self.file_label.setStyleSheet("color: #6c7086;")
        self.file_label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        btn = QPushButton("Browse…")
        btn.clicked.connect(self._browse)
        fgl.addWidget(self.file_label)
        fgl.addWidget(btn)
        layout.addWidget(fg)

        # ── Table ────────────────────────────────────────────────────────────
        tg = QGroupBox("Table")
        tgl = QVBoxLayout(tg)
        self.table_combo = QComboBox()
        self.table_combo.setEnabled(False)
        self.table_info_lbl = QLabel("")
        self.table_info_lbl.setStyleSheet("color: #6c7086; font-size: 11px;")
        self.table_combo.currentTextChanged.connect(self._on_table_changed)
        tgl.addWidget(self.table_combo)
        tgl.addWidget(self.table_info_lbl)
        layout.addWidget(tg)

        layout.addStretch()

    def _browse(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Excel file", "", "Excel Files (*.xlsx *.xlsm)")
        if not path:
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Could not open file:\n{e}")
            return

        table_map = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for tbl in ws.tables.values():
                table_map[tbl.name] = (sheet_name, tbl)

        if not table_map:
            QMessageBox.warning(
                self, "No tables found",
                "No Excel tables (Insert → Table) were found in this workbook.\n"
                "Please format your data as a table in Excel first."
            )
            return

        self.wb = wb
        self._table_map = table_map
        self.file_path = path
        self.file_label.setText(os.path.basename(path))
        self.file_label.setStyleSheet("color: #cdd6f4;")
        self.table_combo.clear()
        self.table_combo.addItems(sorted(table_map.keys()))
        self.table_combo.setEnabled(True)

    def _on_table_changed(self, name):
        if name and name in self._table_map:
            sheet_name, tbl = self._table_map[name]
            self.table_info_lbl.setText(f"Sheet: {sheet_name}   Range: {tbl.ref}")
        else:
            self.table_info_lbl.setText("")

    def get_store(self) -> dict | None:
        """Returns the store config built from the URL, key, and secret fields."""
        url    = self.url_edit.text().strip().rstrip("/")
        key    = self.key_edit.text().strip()
        secret = self.secret_edit.text().strip()
        if not url or not key or not secret:
            return None
        return {"url": url, "consumer_key": key, "consumer_secret": secret}

    def get_table(self):
        """Returns (workbook, sheet_name, table_ref) or None."""
        name = self.table_combo.currentText()
        if not self.wb or not name or name not in self._table_map:
            return None
        sheet_name, tbl = self._table_map[name]
        return self.wb, sheet_name, tbl.ref
